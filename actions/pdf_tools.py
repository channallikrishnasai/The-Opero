"""
pdf_tools.py - Brahma AI PDF support

Creates publication-grade PDF documents from structured content or markdown,
with multi-page running headers/footers, table formatting, chapter pagination,
and deep multi-chapter research synthesis.
"""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

PROJECT_NAME = "Brahma AI - Lite"
DEFAULT_OUTPUT_DIR = Path.home() / "Downloads"


def _sanitize_filename(name: str, default: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._ -]+", "", (name or "").strip())
    safe = re.sub(r"\s+", " ", safe).strip().replace(" ", "_")
    return safe or default


def _resolve_output_path(output_path: str | None, title: str, ext: str, fallback_name: str) -> Path:
    if output_path:
        path = Path(output_path).expanduser()
        if not path.is_absolute():
            head = path.parts[0].lower() if path.parts else ""
            tail = Path(*path.parts[1:]) if len(path.parts) > 1 else Path(path.name)
            if head in {"downloads", "download"}:
                path = Path.home() / "Downloads" / tail
            elif head == "desktop":
                path = Path.home() / "Desktop" / tail
            else:
                path = Path.cwd() / path
        if path.suffix.lower() != ext:
            path = path.with_suffix(ext)
        path.parent.mkdir(parents=True, exist_ok=True)
        return path

    DEFAULT_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return DEFAULT_OUTPUT_DIR / f"{_sanitize_filename(title, fallback_name)}{ext}"


def _open_file(path: Path) -> None:
    try:
        if os.name == "nt":
            subprocess.Popen(["cmd.exe", "/c", "start", "", str(path)], shell=True)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass


def _get_api_key() -> str:
    candidates = [
        Path(__file__).resolve().parent.parent / "config" / "api_keys.json",
        Path.cwd() / "config" / "api_keys.json",
    ]
    for cp in candidates:
        if cp.exists():
            try:
                data = json.loads(cp.read_text(encoding="utf-8"))
                if "gemini_api_key" in data:
                    return data["gemini_api_key"]
            except Exception:
                pass
    return ""


def _parse_json_arg(value, fallback):
    if value is None:
        return fallback
    if isinstance(value, (dict, list)):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return fallback
        try:
            return json.loads(text)
        except Exception:
            return fallback
    return fallback


def _normalize_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return value
    parsed = _parse_json_arg(value, None)
    if isinstance(parsed, list):
        return parsed
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return []
        return [line.strip() for line in re.split(r"\n+", text) if line.strip()]
    return [value]


def _clean_for_reportlab(text: str) -> str:
    """Escapes raw XML entities while preserving legitimate styling tags."""
    if not text:
        return ""
    text = str(text)
    text = re.sub(r"&(?!amp;|lt;|gt;|quot;|apos;|#\d+;)", "&amp;", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*([^*]+)\*", r"<i>\1</i>", text)
    text = re.sub(r"`([^`]+)`", r'<font face="Courier" color="#0F172A"><b>\1</b></font>', text)
    text = text.replace("\u2014", " - ").replace("\u2013", "-")
    return text


def _import_pdf():
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfgen import canvas
        from reportlab.platypus import (
            BaseDocTemplate,
            Frame,
            ListFlowable,
            ListItem,
            PageBreak,
            PageTemplate,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        class NumberedCanvas(canvas.Canvas):
            def __init__(self, *args, **kwargs):
                super().__init__(*args, **kwargs)
                self._saved_page_states = []

            def showPage(self):
                self._saved_page_states.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                num_pages = len(self._saved_page_states)
                for state in self._saved_page_states:
                    self.__dict__.update(state)
                    self.draw_page_decorations(num_pages)
                    super().showPage()
                super().save()

            def draw_page_decorations(self, page_count: int):
                # Suppress headers and footers on cover page (page 1) if document has multiple pages
                if self._pageNumber == 1 and page_count > 1:
                    return
                self.saveState()
                self.setFont("Helvetica", 8)
                self.setFillColor(colors.HexColor("#64748B"))

                # Running header
                doc_title = getattr(self, "doc_title", "Brahma AI Executive Report")
                self.drawString(54, 750, doc_title[:80])
                self.setStrokeColor(colors.HexColor("#CBD5E1"))
                self.setLineWidth(0.5)
                self.line(54, 742, 558, 742)

                # Running footer
                self.line(54, 45, 558, 45)
                self.drawString(54, 32, "BRAHMA AI - AUTONOMOUS INTELLIGENCE RESEARCH")
                page_text = f"Page {self._pageNumber} of {page_count}"
                self.drawRightString(558, 32, page_text)
                self.restoreState()

        return {
            "colors": colors,
            "TA_CENTER": TA_CENTER,
            "TA_JUSTIFY": TA_JUSTIFY,
            "TA_LEFT": TA_LEFT,
            "LETTER": LETTER,
            "ParagraphStyle": ParagraphStyle,
            "getSampleStyleSheet": getSampleStyleSheet,
            "inch": inch,
            "BaseDocTemplate": BaseDocTemplate,
            "Frame": Frame,
            "ListFlowable": ListFlowable,
            "ListItem": ListItem,
            "PageBreak": PageBreak,
            "PageTemplate": PageTemplate,
            "Paragraph": Paragraph,
            "SimpleDocTemplate": SimpleDocTemplate,
            "Spacer": Spacer,
            "Table": Table,
            "TableStyle": TableStyle,
            "NumberedCanvas": NumberedCanvas,
        }
    except Exception as e:
        raise RuntimeError("reportlab is required. Install it with: pip install reportlab") from e


def _parse_markdown_to_blocks(content: str) -> list[dict]:
    lines = content.splitlines()
    blocks = []
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()

        if trimmed.startswith("```"):
            if in_code_block:
                in_code_block = False
                if code_lines:
                    blocks.append({"kind": "code", "text": "\n".join(code_lines)})
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if not trimmed:
            i += 1
            continue

        # Markdown Table Detection
        if trimmed.startswith("|") and trimmed.endswith("|"):
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                row_raw = lines[i].strip()
                if not re.match(r"^\|[\s\-:|]+\|$", row_raw):
                    inner = row_raw[1:-1]
                    cells = [c.strip() for c in inner.split("|")]
                    table_rows.append(" | ".join(cells))
                i += 1
            if table_rows:
                blocks.append({"kind": "table", "text": table_rows})
            continue

        # Headings
        if trimmed.startswith("#"):
            h_match = re.match(r"^(#{1,4})\s*(.*)$", trimmed)
            if h_match:
                level = len(h_match.group(1))
                h_text = h_match.group(2).strip()
                blocks.append({"kind": "heading", "level": level, "text": h_text})
                i += 1
                continue

        # Dividers
        if re.match(r"^(\-{3,}|\*{3,}|_{3,})$", trimmed):
            blocks.append({"kind": "divider"})
            i += 1
            continue

        # Bullets
        bullet_match = re.match(r"^[\*\-\+•]\s+(.*)$", trimmed)
        if bullet_match:
            blocks.append({"kind": "bullet", "text": bullet_match.group(1).strip()})
            i += 1
            continue

        # Numbered list
        num_match = re.match(r"^(\d+)\.\s+(.*)$", trimmed)
        if num_match:
            blocks.append({"kind": "numbered", "text": num_match.group(2).strip()})
            i += 1
            continue

        # Regular paragraph
        para_lines = [trimmed]
        i += 1
        while i < len(lines):
            next_trim = lines[i].strip()
            if not next_trim or next_trim.startswith(("#", "|", "* ", "- ", "+ ", "• ", "```")) or re.match(r"^\d+\.\s+", next_trim) or re.match(r"^(\-{3,}|\*{3,}|_{3,})$", next_trim):
                break
            para_lines.append(next_trim)
            i += 1
        blocks.append({"kind": "body", "text": " ".join(para_lines)})

    return blocks


def _open_docx_text(path: Path) -> list[dict]:
    from docx import Document

    doc = Document(path)
    blocks: list[dict] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        kind = "body"
        level = 0
        if style_name.startswith("Heading "):
            kind = "heading"
            try:
                level = int(style_name.split()[-1])
            except Exception:
                level = 1
        blocks.append({"kind": kind, "level": level, "text": text})

    for table in doc.tables:
        rows = []
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                rows.append(row_text)
        if rows:
            blocks.append({"kind": "table", "text": rows})

    return blocks


def _plain_text_blocks(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    if any(line.strip().startswith(("#", "|", "* ", "- ")) for line in text.splitlines()[:20]):
        return _parse_markdown_to_blocks(text)
    blocks = []
    for chunk in [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]:
        blocks.append({"kind": "body", "text": chunk})
    return blocks


def _structured_blocks(parameters: dict) -> list[dict]:
    content = parameters.get("content") or parameters.get("body") or ""

    if content and any(pattern in str(content) for pattern in ["\n#", "# ", "\n|", "\n* ", "\n- "]):
        return _parse_markdown_to_blocks(str(content))

    sections = _parse_json_arg(parameters.get("sections"), None)
    if sections:
        if isinstance(sections, dict):
            sections = [sections]
        blocks = []
        for section in sections:
            if not isinstance(section, dict):
                continue
            heading = (section.get("heading") or section.get("title") or "").strip()
            body = section.get("body") or section.get("content") or ""
            bullets = section.get("bullets") or []
            level = int(section.get("level") or 1)
            if heading:
                blocks.append({"kind": "heading", "level": level, "text": heading})
            if isinstance(body, list):
                for item in body:
                    blocks.append({"kind": "body", "text": str(item)})
            else:
                for para in [p.strip() for p in re.split(r"\n\s*\n", str(body)) if p.strip()]:
                    blocks.append({"kind": "body", "text": para})
            for bullet in _normalize_list(bullets):
                blocks.append({"kind": "bullet", "text": str(bullet)})
        return blocks

    paragraphs = _normalize_list(parameters.get("paragraphs"))
    bullets = _normalize_list(parameters.get("bullets"))
    numbered = _normalize_list(parameters.get("numbered"))

    blocks = []
    if paragraphs:
        for para in paragraphs:
            blocks.append({"kind": "body", "text": str(para)})
    elif content:
        for para in [p.strip() for p in re.split(r"\n\s*\n", str(content)) if p.strip()]:
            blocks.append({"kind": "body", "text": para})

    for bullet in bullets:
        blocks.append({"kind": "bullet", "text": str(bullet)})
    for item in numbered:
        blocks.append({"kind": "numbered", "text": str(item)})

    return blocks


def synthesize_deep_report(goal_or_topic: str, title: str, research_notes: str = "") -> str:
    """
    Synthesizes an exhaustive, 10-chapter technical research monograph
    using high-speed Gemini-3.1-flash-lite in two cohesive passes.
    """
    api_key = _get_api_key()
    if not api_key:
        return ""

    import google.generativeai as genai
    genai.configure(api_key=api_key)

    model_names = ["gemini-3.1-flash-lite", "gemini-3.5-flash", "gemini-flash-latest"]
    model = None
    for name in model_names:
        try:
            model = genai.GenerativeModel(name)
            break
        except Exception:
            continue

    if not model:
        return ""

    print(f"[PDF Tools] [*] Synthesizing exhaustive research monograph on: {title}...")

    prompts = [
        f"""You are a Distinguished AI Scientist and Executive Technical Author.
The user requested: '{goal_or_topic}'
Title: '{title}'
Research data gathered:
{research_notes[:3000]}

Write PART 1 of an exhaustive, publication-grade research monograph in Markdown.
Cover Chapters 1 to 5 in full technical depth:
# Chapter 1: Executive Summary & Foundational Overview
## Conceptual Framework and Definitions
## Historical Context and Technological Genesis
## High-Level Paradigm Shifts

# Chapter 2: Core Architecture & Theoretical Foundations
## Primary Operational Cycles and Feedback Loops
## Sensory Ingestion and State Representation
## Deliberative Compute and Latency Trade-offs

# Chapter 3: Behavioral Typologies & Functional Taxonomies
## Classification Models and Structural Patterns
## Goal Formulation and Utility Maximization
## Adaptive Learning and Iterative Convergence

# Chapter 4: Algorithmic Mechanics & Decision Engines
## Deliberative Reasoning Frameworks (e.g. ReAct, Tree Search, Reflexion)
## Prompt Optimization and Context Distillation
## Error Recovery, Exception Trapping, and Self-Healing

# Chapter 5: Retrieval Dynamics, Memory Architectures & Tool Use
## Epistemic Short-Term vs. Parametric Long-Term Storage
## Dense Vector Indexing and Hybrid Retrieval Integration
## Sandboxed Tool Interfaces, RPC Protocols and Safety Enclaves

Write substantial, deeply technical paragraphs under every subsection (300+ words per section).
Include detailed bullet points and formal explanations.
""",
        f"""You are a Distinguished AI Scientist and Executive Technical Author.
Write PART 2 of an exhaustive, publication-grade research monograph on '{title}' in Markdown.
Cover Chapters 6 to 10 in full technical depth:
# Chapter 6: Multi-Component Collaboration & Swarm Intelligence
## Topology Design: Hierarchical vs. Decentralized Swarms
## Inter-Node Communication Protocols and Message Passing
## Consensus Algorithms and Conflict Resolution

# Chapter 7: Enterprise Applications & Real-World Deployments
## Production Engineering and Automated System Maintenance
## High-Throughput Quantitative and Analytical Pipelines
## Critical Infrastructure, Healthcare, and Cyberdefense Deployments

# Chapter 8: Comparative Architecture Matrix
## Performance and Latency Trade-off Analysis
(Include a detailed markdown table comparing key paradigms across Latency, Complexity, Autonomy, Reliability, Cost)
## Framework and Ecosystem Benchmark Matrix
(Include a detailed markdown table comparing modern production frameworks across Scalability, State Management, Developer Overhead)

# Chapter 9: Robustness, Security Guardrails & Alignment
## Adversarial Vectors, Prompt Injections and Jailbreaks
## Hallucination Mitigation, Goal Drift, and Runaway Recursion
## Sandboxing, Least Privilege Architecture, and Human-in-the-Loop Safeguards

# Chapter 10: Future Horizons, Technological Roadmap & Conclusion
## Long-Term Evolution Toward General Problem Solvers
## Socio-Technical and Economic Implications
## Concluding Technical Synthesis

Write substantial, deeply technical paragraphs under every subsection. Ensure all markdown tables are complete with clean columns.
"""
    ]

    parts = []
    for idx, prompt_text in enumerate(prompts):
        try:
            t0 = time.time()
            resp = model.generate_content(prompt_text)
            if resp.text:
                parts.append(resp.text)
                print(f"[PDF Tools] [+] Monograph Part {idx+1} generated in {time.time()-t0:.1f}s ({len(resp.text.split())} words)")
        except Exception as e:
            print(f"[PDF Tools] [!] Part {idx+1} generation failed: {e}")

    return "\n\n".join(parts)


def _render_title_page(story, pdf, title: str, subtitle: str | None, styles, author: str | None = None):
    story.append(pdf["Spacer"](1, 1.8 * pdf["inch"]))
    story.append(pdf["Paragraph"]("TECHNICAL RESEARCH MONOGRAPH", styles["brahma_doc_badge"]))
    story.append(pdf["Paragraph"](title, styles["brahma_doc_title"]))
    if subtitle:
        story.append(pdf["Paragraph"](subtitle, styles["brahma_doc_subtitle"]))
    story.append(pdf["Spacer"](1, 0.4 * pdf["inch"]))
    author_text = author or "Suryaansh Tiwari"
    story.append(
        pdf["Paragraph"](
            f"<b>Author:</b> {author_text} &nbsp;|&nbsp; <b>Division:</b> Brahma AI Autonomous Systems &nbsp;|&nbsp; <b>Date:</b> {datetime.now().strftime('%B %d, %Y')}",
            styles["brahma_doc_meta"],
        )
    )
    story.append(pdf["Spacer"](1, 0.1 * pdf["inch"]))
    story.append(
        pdf["Paragraph"](
            "<b>Classification:</b> Public Technical Monograph &nbsp;|&nbsp; <b>Engine:</b> Brahma AI Publication Core",
            styles["brahma_doc_meta"],
        )
    )
    story.append(pdf["PageBreak"]())


def _pdf_story_from_blocks(blocks: list[dict], pdf, styles):
    story = []
    first_heading = True

    for block in blocks:
        kind = block.get("kind", "body")
        text = str(block.get("text", "")).strip()
        if not text and kind != "divider":
            continue

        if kind == "heading":
            level = int(block.get("level") or 1)
            clean_t = _clean_for_reportlab(text)
            if level == 1:
                if not first_heading:
                    story.append(pdf["PageBreak"]())
                first_heading = False
                story.append(pdf["Paragraph"](clean_t, styles["brahma_h1"]))
                story.append(pdf["Spacer"](1, 0.08 * pdf["inch"]))
            elif level == 2:
                story.append(pdf["Paragraph"](clean_t, styles["brahma_h2"]))
                story.append(pdf["Spacer"](1, 0.05 * pdf["inch"]))
            else:
                story.append(pdf["Paragraph"](clean_t, styles["brahma_h3"]))
                story.append(pdf["Spacer"](1, 0.04 * pdf["inch"]))

        elif kind == "bullet":
            clean_t = _clean_for_reportlab(text)
            story.append(pdf["Paragraph"](f"&bull;&nbsp;&nbsp;{clean_t}", styles["brahma_bullet"]))

        elif kind == "numbered":
            clean_t = _clean_for_reportlab(text)
            story.append(pdf["Paragraph"](f"<b>{clean_t}</b>", styles["brahma_bullet"]))

        elif kind == "code":
            clean_t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(pdf["Paragraph"](f"<pre>{clean_t}</pre>", styles["brahma_code"]))

        elif kind == "divider":
            story.append(pdf["Spacer"](1, 0.1 * pdf["inch"]))

        elif kind == "table":
            rows = block.get("text") or []
            if not rows:
                continue
            parsed_rows = []
            for r_idx, row in enumerate(rows):
                cells = row.split(" | ")
                row_flowables = []
                for cell in cells:
                    c_clean = _clean_for_reportlab(cell)
                    if r_idx == 0:
                        row_flowables.append(pdf["Paragraph"](c_clean, styles["brahma_table_header"]))
                    else:
                        row_flowables.append(pdf["Paragraph"](c_clean, styles["brahma_table_cell"]))
                parsed_rows.append(row_flowables)

            col_count = len(parsed_rows[0])
            avail_width = 504.0
            col_w = avail_width / max(col_count, 1)

            t = pdf["Table"](parsed_rows, colWidths=[col_w] * col_count, repeatRows=1)
            t_style = [
                ("BACKGROUND", (0, 0), (-1, 0), pdf["colors"].HexColor("#0F172A")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("GRID", (0, 0), (-1, -1), 0.5, pdf["colors"].HexColor("#CBD5E1")),
            ]
            for r in range(1, len(parsed_rows)):
                bg = pdf["colors"].HexColor("#F8FAFC") if r % 2 == 1 else pdf["colors"].white
                t_style.append(("BACKGROUND", (0, r), (-1, r), bg))

            t.setStyle(pdf["TableStyle"](t_style))
            story.append(pdf["Spacer"](1, 0.08 * pdf["inch"]))
            story.append(t)
            story.append(pdf["Spacer"](1, 0.12 * pdf["inch"]))

        else:
            clean_t = _clean_for_reportlab(text)
            story.append(pdf["Paragraph"](clean_t, styles["brahma_body"]))

    return story


def create_pdf(parameters: dict, player=None) -> str:
    pdf = _import_pdf()
    title = (parameters.get("title") or parameters.get("name") or "Document").strip()
    subtitle = (parameters.get("subtitle") or "").strip()
    output_path = _resolve_output_path(parameters.get("output_path"), title, ".pdf", "brahma_ai_output")
    auto_open = parameters.get("auto_open", True)
    action = (parameters.get("action") or "create").lower().strip()
    source_path_str = (parameters.get("file_path") or "").strip()
    source_path = Path(source_path_str) if source_path_str else None

    # Check if a deep research synthesis is warranted
    raw_content = str(parameters.get("content") or parameters.get("body") or "")
    goal_hint = str(parameters.get("goal") or title)
    wants_report = any(w in (goal_hint + " " + title).lower() for w in ["research", "report", "monograph", "comprehensive", "10-20 page", "10 page", "paper", "deep dive"])

    # Only synthesize if content is brief (< 400 chars) AND does not already have multiple chapters
    if wants_report and (len(raw_content) < 400 and "# Chapter" not in raw_content or action == "deep_report"):
        synth = synthesize_deep_report(goal_hint, title, research_notes=raw_content)
        if synth:
            parameters["content"] = synth
            if not subtitle:
                subtitle = "Comprehensive Technical Analysis and Operational Architecture"

    styles = pdf["getSampleStyleSheet"]()

    styles.add(pdf["ParagraphStyle"](
        name="brahma_doc_badge",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=12,
        textColor=pdf["colors"].HexColor("#2563EB"),
        alignment=pdf["TA_CENTER"],
        spaceAfter=15,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_doc_title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=24,
        leading=30,
        textColor=pdf["colors"].HexColor("#0B192C"),
        alignment=pdf["TA_CENTER"],
        spaceAfter=12,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_doc_subtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=12,
        leading=16,
        textColor=pdf["colors"].HexColor("#334155"),
        alignment=pdf["TA_CENTER"],
        spaceAfter=20,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_doc_meta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=pdf["colors"].HexColor("#64748B"),
        alignment=pdf["TA_CENTER"],
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_h1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15.5,
        leading=19,
        textColor=pdf["colors"].HexColor("#0F172A"),
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_h2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12.5,
        leading=15.5,
        textColor=pdf["colors"].HexColor("#1E293B"),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_h3",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13.5,
        textColor=pdf["colors"].HexColor("#334155"),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        textColor=pdf["colors"].HexColor("#1E293B"),
        alignment=pdf["TA_JUSTIFY"],
        spaceAfter=7,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_bullet",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=pdf["colors"].HexColor("#1E293B"),
        leftIndent=14,
        spaceAfter=4,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_table_header",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=pdf["colors"].white,
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_table_cell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10.5,
        textColor=pdf["colors"].HexColor("#0F172A"),
    ))

    styles.add(pdf["ParagraphStyle"](
        name="brahma_code",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=8,
        leading=10,
        textColor=pdf["colors"].HexColor("#0F172A"),
        backColor=pdf["colors"].HexColor("#F1F5F9"),
        borderPadding=5,
        spaceBefore=5,
        spaceAfter=7,
    ))

    blocks = []
    if action == "convert" and source_path:
        if not source_path.exists():
            return f"File not found: {source_path}"
        suffix = source_path.suffix.lower()
        if suffix == ".docx":
            blocks = _open_docx_text(source_path)
            if not blocks:
                return "The DOCX appears to be empty."
        elif suffix in {".txt", ".md", ".rst", ".log"}:
            blocks = _plain_text_blocks(source_path)
        else:
            return f"Unsupported source type for PDF conversion: {source_path.suffix}"
    else:
        blocks = _structured_blocks(parameters)

    if not blocks:
        blocks = [{"kind": "body", "text": "Document content ready."}]

    doc = pdf["SimpleDocTemplate"](
        str(output_path),
        pagesize=pdf["LETTER"],
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
        title=title,
        author=parameters.get("author") or "Suryaansh Tiwari",
        subject=parameters.get("subject") or "Research Document",
    )
    doc.doc_title = title

    story = []
    has_chapters = any(b.get("kind") == "heading" and b.get("level") == 1 for b in blocks)
    if title and has_chapters:
        _render_title_page(story, pdf, title, subtitle, styles, author=parameters.get("author"))

    if action == "create_letter":
        recipient = (parameters.get("recipient") or parameters.get("to") or "").strip()
        date_value = (parameters.get("date") or datetime.now().strftime("%B %d, %Y")).strip()
        salutation = (parameters.get("salutation") or (f"Dear {recipient}," if recipient else "Dear Sir or Madam,")).strip()
        closing = (parameters.get("closing") or "Sincerely,").strip()
        signature = (parameters.get("signature") or parameters.get("author") or "Suryaansh Tiwari").strip()
        body_text = parameters.get("body") or parameters.get("content") or ""

        story.append(pdf["Paragraph"](date_value, styles["brahma_body"]))
        story.append(pdf["Spacer"](1, 0.08 * pdf["inch"]))
        if recipient:
            story.append(pdf["Paragraph"](recipient, styles["brahma_body"]))
            story.append(pdf["Spacer"](1, 0.06 * pdf["inch"]))
        story.append(pdf["Paragraph"](salutation, styles["brahma_body"]))
        story.append(pdf["Spacer"](1, 0.08 * pdf["inch"]))
        for para in _normalize_list(parameters.get("paragraphs")) or [p.strip() for p in re.split(r"\n\s*\n", str(body_text)) if p.strip()]:
            story.append(pdf["Paragraph"](str(para), styles["brahma_body"]))
            story.append(pdf["Spacer"](1, 0.08 * pdf["inch"]))
        story.append(pdf["Spacer"](1, 0.18 * pdf["inch"]))
        story.append(pdf["Paragraph"](closing, styles["brahma_body"]))
        story.append(pdf["Spacer"](1, 0.3 * pdf["inch"]))
        story.append(pdf["Paragraph"](signature, styles["brahma_body"]))
    else:
        story.extend(_pdf_story_from_blocks(blocks, pdf, styles))

    try:
        doc.build(story, canvasmaker=pdf["NumberedCanvas"])
    except Exception as e:
        return f"PDF creation failed: {e}"

    if auto_open:
        _open_file(output_path)
    return f"PDF created: {output_path}"
